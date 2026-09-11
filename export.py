"""Write-only Excel export: complete raw data, lineage and safe text cells."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import math
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sped_parser import get_config, conflicting_files


def export_workbook(batch, documents=None, filtered=False, row_limit=None):
    """Filtered exports contain only visible documents; full exports retain every record."""
    row_limit = min(row_limit or get_config("export.max_rows_per_sheet", 1048575), 1048575)
    if row_limit < 1:
        raise ValueError("Limite de linhas inválido.")
    tables = {
        "LEIA_ME": pd.DataFrame(
            {
                "INFORMACAO": [
                    "Extrator SPED v6 — dados de conferência; não substitui o PVA.",
                    "Escopo: documentos filtrados."
                    if filtered
                    else "Escopo: lote completo, registros aceitos, originais e ocorrências.",
                    "Cada documento aparece uma vez. Tributos são os valores informados nos cabeçalhos; não são uma apuração.",
                    "Registros originais preservam texto. Linhas rejeitadas/não suportadas ficam fora dos indicadores.",
                    "REGISTRO_ID identifica arquivo e linha; PAI_ID aponta para o registro imediatamente superior mapeado.",
                    "Campos com mais de 32.767 caracteres são divididos em colunas __PARTE_N; concatene para reconstruir.",
                    "Caracteres de controle incompatíveis com XML são escritos como escapes \\uNNNN nas células.",
                    "O Excel possui precisão numérica limitada; as abas ORIG preservam os números exatamente como recebidos.",
                    "CFOP seleciona documentos que contêm o código. O valor exibido é o documento inteiro, não um rateio por CFOP.",
                ]
            }
        )
    }
    tables["DOCUMENTOS"] = batch.documents if documents is None else documents
    if not filtered:
        tables["ARQUIVOS"] = batch.manifest
        tables["OCORRENCIAS"] = batch.issues
        tables["ARQUIVOS_REJEITADOS"] = pd.DataFrame(batch.rejected_files)
        tables["SOBREPOSICOES"] = pd.DataFrame(
            conflicting_files(batch), columns=["ARQUIVO", "SOBREPOSTO_A"]
        )
        for attr, prefix in [("tables", "REG"), ("raw_tables", "ORIG")]:
            codes = sorted({k for f in batch.files for k in getattr(f, attr)})
            for code in codes:
                tables[f"{prefix}_{code}"] = pd.concat(
                    [getattr(f, attr)[code] for f in batch.files if code in getattr(f, attr)],
                    ignore_index=True,
                )
    estimated = sum(max(1, len(df)) * max(1, len(df.columns)) for df in tables.values())
    if estimated > get_config("export.max_cells", 5000000):
        raise ValueError(
            "Exportação excede o limite de células. Exporte uma seleção menor ou processe menos arquivos."
        )
    book = Workbook(write_only=True)
    names = set()
    for name, df in tables.items():
        if df.empty:
            df = pd.DataFrame({"INFORMACAO": ["Nenhum registro neste escopo."]})
        # Expand long string columns before writing to avoid Excel's silent truncation.
        specs = []
        for col in df.columns:
            longest = max((len(str(v)) for v in df[col] if isinstance(v, str)), default=0)
            chunks = max(1, math.ceil(longest / 32767))
            for part in range(chunks):
                specs.append((col, part, chunks, f"{col}__PARTE_{part + 1}" if chunks > 1 else col))
        if len(specs) > 16384:
            raise ValueError("Quantidade de colunas excede o limite do Excel.")
        for page, start in enumerate(range(0, len(df), row_limit), 1):
            suffix = f"_{page}" if len(df) > row_limit else ""
            safe = re.sub(r"[\\/*?:\[\]]", "_", name)[: 31 - len(suffix)] + suffix
            base = safe
            n = 1
            while safe in names:
                n += 1
                safe = base[:26] + f"_{n}"
            names.add(safe)
            sheet = book.create_sheet(safe)
            sheet.freeze_panes = "A2"
            header = []
            for _, _, _, label in specs:
                cell = WriteOnlyCell(sheet, value=label)
                cell.data_type = "s"
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="173F46")
                header.append(cell)
            sheet.append(header)
            for values in df.iloc[start : start + row_limit].to_dict("records"):
                cells = []
                for col, part, chunks, _ in specs:
                    value = values[col]
                    if value is None or (not isinstance(value, (list, dict)) and pd.isna(value)):
                        value = None
                    if isinstance(value, str):
                        value = value[part * 32767 : (part + 1) * 32767] if chunks > 1 else value
                        value = ILLEGAL_CHARACTERS_RE.sub(
                            lambda m: f"\\u{ord(m.group()):04x}", value
                        )
                        # XML escaping can expand a string; do not truncate it silently.
                        if len(value) > 32767:
                            raise ValueError(
                                "Campo excede o limite após escapar caracteres de controle. Revise o arquivo de origem."
                            )
                    cell = WriteOnlyCell(sheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    elif isinstance(value, (date, datetime)):
                        cell.number_format = "dd/mm/yyyy"
                    elif isinstance(value, Decimal):
                        cell.number_format = "#,##0.00########"
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(specs))}{min(row_limit, len(df) - start) + 1}"
            )
    output = BytesIO()
    book.save(output)
    return output.getvalue()
