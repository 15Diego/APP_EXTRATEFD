from datetime import date
from decimal import Decimal
from io import BytesIO
import pandas as pd
import pytest
from openpyxl import load_workbook
from demo import record, sample_file
from sped_parser import (
    process_batch,
    process_content,
    filter_documents,
    calculate_totals,
    conflicting_files,
    parse_number,
    detect_efd_type,
    CONFIG,
)
from export import export_workbook
from validators import validate_cnpj


def batch(kind="ICMS_IPI", count=1):
    return process_batch([("example.txt", sample_file(kind, count))])


def insert(content, line):
    lines = content.decode().splitlines(keepends=True)
    lines.insert(-1, line)
    lines[-1] = record("ICMS_IPI", "9999", QTD_LIN=str(len(lines)))
    return "".join(lines).encode()


def test_document_not_multiplied_by_children():
    b = batch()
    assert len(b.files[0].tables["C170"]) == 2
    assert len(b.files[0].tables["C190"]) == 2
    totals = calculate_totals(b.documents)
    assert totals["documents"] == 1
    assert totals["value"] == Decimal("100")
    assert totals["icms"] == Decimal("18")
    assert b.issues.empty


@pytest.mark.parametrize("kind", ["ICMS_IPI", "CONTRIBUICOES"])
def test_detection_without_optional_blocks(kind):
    content = sample_file(kind)
    assert detect_efd_type(content) == kind
    assert len(batch(kind).documents) == 1


def test_operation_cfop_date_and_literal_filters():
    b = batch(count=6)
    assert len(filter_documents(b, {"operations": ["Entrada"]})) == 3
    assert len(filter_documents(b, {"operations": ["Saída"]})) == 3
    assert len(filter_documents(b, {"cfops": ["9999"]})) == 0
    assert len(filter_documents(b, {"cfops": ["2102"]})) == 3
    assert len(filter_documents(b, {"start": date(2026, 1, 11), "end": date(2026, 1, 12)})) == 2
    assert len(filter_documents(b, {"search": "["})) == 0
    assert len(filter_documents(b, {"search": "11222333000181"})) == 6


def test_decimal_precision_and_fractional_quantity():
    assert parse_number("1.234,56789") == Decimal("1234.56789")
    assert parse_number("0,1") + parse_number("0,2") == Decimal("0.3")
    with pytest.raises(ValueError):
        parse_number("12.34")
    with pytest.raises(ValueError):
        parse_number("NaN")


def test_quarantine_unknown_and_empty_not_success():
    content = insert(insert(sample_file(count=1), "|C100|\n"), "|ZZZZ|preserved|\n")
    result = process_content(content, "bad.txt")
    assert result.metadata["REJEITADAS"] == 1
    assert result.metadata["NAO_SUPORTADAS"] == 1
    assert result.metadata["APROVEITAMENTO_PCT"] < 100
    assert result.raw_tables["ZZZZ"].iloc[0]["CAMPO_001"] == "preserved"
    with pytest.raises(ValueError, match="estrito"):
        process_content(content, "bad.txt", strict=True)


def test_field_count_no_truncation():
    content = insert(sample_file(count=1), "|0150|P2|" + ("|".join([""] * 11)) + "|EXTRA|\n")
    r = process_content(content, "extra.txt")
    assert "EXTRA" in r.raw_tables["0150"].iloc[-1]["CONTEUDO_ORIGINAL"]
    assert r.metadata["REJEITADAS"] >= 1


def test_invalid_dates_rejected_with_origin():
    r = process_content(sample_file(count=1).replace(b"10012026", b"32012026"), "date.txt")
    assert r.documents.empty
    assert any(i["CAMPO"] == "DT_DOC" and i["LINHA"] > 0 for i in r.issues)


def test_duplicate_files_and_partial_batch():
    content = sample_file(count=1)
    b = process_batch([("a.txt", content), ("copy.txt", content), ("bad.txt", b"wrong")])
    assert len(b.files) == 1
    assert len(b.failures) == 2
    assert calculate_totals(b.documents)["documents"] == 1


def test_overlaps_and_unique_lineage():
    content = sample_file(count=1)
    b = process_batch([("a.txt", content), ("b.txt", content.replace(b"1000", b"9000"))])
    assert len(b.files) == 2
    assert len(set(b.documents["DOCUMENTO_ID"])) == 2
    assert conflicting_files(b)
    assert not conflicting_files(b, [b.files[0].metadata["ARQUIVO_ID"]])


def test_wrong_type_and_size_rejected(monkeypatch):
    with pytest.raises(ValueError, match="tipo"):
        process_content(sample_file(), "wrong.txt", expected_type="CONTRIBUICOES")
    monkeypatch.setitem(CONFIG["processing"], "max_file_size_mb", 0)
    with pytest.raises(ValueError, match="maior"):
        process_content(sample_file(), "large.txt")


def test_m200_m500_and_inventory_numbers_exported():
    content = sample_file("CONTRIBUICOES", 1)
    text = content.decode()
    extra = (
        record("CONTRIBUICOES", "M001", IND_MOV="0")
        + record("CONTRIBUICOES", "M200", VL_TOT_CONT_REC="10,00")
        + record("CONTRIBUICOES", "M500", VL_CRED="23,456")
    )
    lines = text.splitlines(keepends=True)
    lines[-1:-1] = extra.splitlines(keepends=True)
    lines[-1] = record("CONTRIBUICOES", "9999", QTD_LIN=str(len(lines)))
    b = process_batch([("contrib.txt", "".join(lines).encode())])
    assert b.files[0].tables["M500"]["VL_CRED"].iloc[0] == Decimal("23.456")
    wb = load_workbook(BytesIO(export_workbook(b)))
    assert "REG_M200" in wb.sheetnames and "ORIG_M500" in wb.sheetnames


def test_excel_empty_formula_long_text_and_pagination():
    b = batch()
    b.files[0].raw_tables["TEST"] = pd.DataFrame({"TEXT": ["=1+1", "a" * 40000, "[x]"]})
    wb = load_workbook(BytesIO(export_workbook(b, row_limit=2)))
    assert wb["ORIG_TEST_1"]["A2"].data_type == "s"
    assert wb["ORIG_TEST_1"]["A2"].value == "=1+1"
    assert len(wb["ORIG_TEST_1"]["A3"].value + wb["ORIG_TEST_1"]["B3"].value) == 40000
    assert "ORIG_TEST_2" in wb.sheetnames
    empty = load_workbook(BytesIO(export_workbook(b, b.documents.iloc[:0], filtered=True)))
    assert "DOCUMENTOS" in empty.sheetnames
    assert empty["DOCUMENTOS"]["A2"].value.startswith("Nenhum")


def test_cancelled_excluded_from_kpis():
    content = sample_file(count=1).replace(b"|55|00|", b"|55|02|")
    b = process_batch([("cancel.txt", content)])
    assert calculate_totals(b.documents)["value"] == 0
    assert filter_documents(b, {}).empty
    assert len(filter_documents(b, {"include_cancelled": True})) == 1


@pytest.mark.parametrize(
    "cnpj", ["11222333000181", "11.222.333/0001-81", "12.ABC.345/01DE-35", "00.000.000/E08G-12"]
)
def test_numeric_and_alphanumeric_cnpj(cnpj):
    assert validate_cnpj(cnpj)


@pytest.mark.parametrize(
    "cnpj", ["00000000000000", "11222333000180", "xx11222333000181", "12ABC34501DE00"]
)
def test_invalid_cnpj(cnpj):
    assert not validate_cnpj(cnpj)


def test_nested_parent_and_reset_after_bad_parent():
    lines = sample_file(count=1).decode().splitlines(keepends=True)
    insert_at = next(i for i, line in enumerate(lines) if line.startswith("|C190|"))
    lines[insert_at:insert_at] = [
        record("ICMS_IPI", "C195", COD_OBS="1"),
        record("ICMS_IPI", "C197", COD_AJ="AJ", VL_ICMS="1,00"),
        "|C195|\n",
        record("ICMS_IPI", "C197", COD_AJ="BAD", VL_ICMS="2,00"),
    ]
    lines[-1] = record("ICMS_IPI", "9999", QTD_LIN=str(len(lines)))
    r = process_content("".join(lines).encode(), "nested.txt")
    assert len(r.tables["C197"]) == 1
    assert r.tables["C197"].iloc[0]["PAI_ID"] == r.tables["C195"].iloc[0]["REGISTRO_ID"]
    assert any(i["CAMPO"] == "VINCULO" for i in r.issues)


def test_unknown_header_rejected():
    with pytest.raises(ValueError):
        process_content(b"|C100|\n", "bad.txt")


def test_cross_total_discrepancy_visible():
    content = sample_file(count=1).replace(b"|50,00|", b"|49,00|")
    r = process_content(content, "totals.txt")
    assert any(i["CAMPO"] == "VL_MERC" for i in r.issues)


def test_current_fiscal_catalog_against_official_positions():
    from schema import get_layout_config, parents_for

    fields = get_layout_config("ICMS_IPI")[0]
    # Independent expectations from the record tables in Guia Prático 3.2.2.
    assert len(fields["C170"]) == 38 and fields["C170"][37] == "VL_ABAT_NT"
    assert len(fields["C176"]) == 27 and fields["C176"][26] == "VL_UNIT_RES_FCP_ST"
    assert len(fields["C500"]) == 40 and fields["C500"][39] == "OUTRAS_DED"
    assert fields["K200"] == ["REG", "DT_EST", "COD_ITEM", "QTD", "IND_EST", "COD_PART"]
    assert fields["0221"] == ["REG", "COD_ITEM_ATOMICO", "QTD_CONTIDA"]
    assert len(fields["D700"]) == 32 and fields["D700"][21] == "CHV_DOCe"
    assert parents_for("ICMS_IPI")["C181"] == "C170"
    assert parents_for("ICMS_IPI")["C186"] == "C100"
    assert parents_for("ICMS_IPI")["D731"] == "D730"


def test_rejected_file_content_is_recoverable():
    import base64

    content = b"invalid\x00\xff"
    b = process_batch([("broken.txt", content)])
    assert base64.b64decode(b.rejected_files[0]["CONTEUDO_BASE64"]) == content
    wb = load_workbook(BytesIO(export_workbook(b)))
    assert "ARQUIVOS_REJEITADOS" in wb.sheetnames


def test_formatted_cnpj_search_and_mixed_efd_totals():
    b = process_batch(
        [
            ("fiscal.txt", sample_file("ICMS_IPI", 1)),
            ("contrib.txt", sample_file("CONTRIBUICOES", 1)),
        ]
    )
    assert len(b.files) == 2
    fiscal = filter_documents(b, {"type": "ICMS_IPI", "search": "11.222.333/0001-81"})
    assert calculate_totals(fiscal)["value"] == 100
    assert len(fiscal) == 1


def test_orphan_after_sibling_does_not_link_to_previous_item():
    content = sample_file(count=1)
    lines = content.decode().splitlines(keepends=True)
    pos = next(i for i, line in enumerate(lines) if line.startswith("|C990|"))
    lines.insert(pos, record("ICMS_IPI", "C171", NUM_TANQUE="1", QTDE="2,00"))
    lines[-1] = record("ICMS_IPI", "9999", QTD_LIN=str(len(lines)))
    r = process_content("".join(lines).encode(), "orphan.txt")
    assert "C171" not in r.tables
    assert r.raw_tables["C171"]["STATUS"].iloc[0] == "Rejeitado"


def test_nfcom_and_inventory_are_separate_from_invoice_details():
    lines = sample_file(count=1).decode().splitlines(keepends=True)
    additions = [
        record("ICMS_IPI", "D001", IND_MOV="0"),
        record(
            "ICMS_IPI",
            "D700",
            IND_OPER="0",
            IND_EMIT="1",
            COD_PART="P1",
            COD_MOD="62",
            COD_SIT="00",
            NUM_DOC="8000",
            DT_DOC="20012026",
            VL_DOC="50,00",
        ),
        record("ICMS_IPI", "D730", CFOP="1303", VL_OPR="50,00"),
        record("ICMS_IPI", "D731", VL_FCP_OP="1,00"),
        record("ICMS_IPI", "H001", IND_MOV="0"),
        record("ICMS_IPI", "H005", DT_INV="31012026", VL_INV="100,00"),
        record("ICMS_IPI", "H010", COD_ITEM="ITEM1", QTD="1,50", VL_UNIT="2,00", VL_ITEM="3,00"),
    ]
    lines[-1:-1] = additions
    lines[-1] = record("ICMS_IPI", "9999", QTD_LIN=str(len(lines)))
    r = process_content("".join(lines).encode(), "extended.txt")
    assert calculate_totals(r.documents)["value"] == 150
    assert calculate_totals(r.documents)["documents"] == 2
    assert r.tables["H010"]["QTD"].iloc[0] == Decimal("1.50")
    assert r.tables["D731"]["PAI_ID"].iloc[0] == r.tables["D730"]["REGISTRO_ID"].iloc[0]
